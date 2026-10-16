"""Build the editable VBot deck; export to PDF with PowerPoint on Windows."""

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent
PPTX = OUT / "VBot_Eviaura_DRAFT.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(17, 35, 54)
INK = RGBColor(31, 48, 64)
TEAL = RGBColor(0, 113, 119)
CYAN = RGBColor(220, 242, 239)
ORANGE = RGBColor(224, 114, 53)
PALE = RGBColor(252, 245, 238)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(248, 250, 249)
MUTED = RGBColor(92, 110, 118)
LINE = RGBColor(215, 225, 226)
RED = RGBColor(171, 63, 48)


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.7)
    if radius:
        try:
            shape.adjustments[0] = 0.1
        except Exception:
            pass
    return shape


def txt(slide, text, x, y, w, h, size=18, color=INK, bold=False,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Arial"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    for i, part in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = part
        p.alignment = align
        p.space_after = Pt(4)
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def pill(slide, label, x, y, w, color=TEAL):
    box(slide, x, y, w, 0.32, color, color)
    txt(slide, label, x + 0.08, y + 0.035, w - 0.16, 0.25, 10, WHITE, True)


def base(kicker, title, subtitle=None, page=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    box(slide, 0, 0, 13.333, 0.09, TEAL, TEAL, False)
    txt(slide, kicker.upper(), 0.58, 0.24, 11.8, 0.27, 10, TEAL, True)
    txt(slide, title, 0.58, 0.66, 12.2, 0.64, 29, NAVY, True)
    if subtitle:
        txt(slide, subtitle, 0.6, 1.32, 12.0, 0.47, 13, MUTED)
    box(slide, 0.58, 7.05, 12.16, 0.01, LINE, LINE, False)
    txt(slide, "EVIAURA  /  VBOT PRODUCTION CAPSTONE", 0.59, 7.12, 8, 0.19, 8.5, MUTED, True)
    txt(slide, f"{page or len(prs.slides):02d} / 12", 11.8, 7.1, 0.9, 0.2, 8.5, MUTED, True, PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, label, body, accent=TEAL, body_size=14):
    box(slide, x, y, w, h)
    box(slide, x, y, 0.07, h, accent, accent, False)
    txt(slide, label, x + 0.22, y + 0.18, w - 0.42, 0.38, 16, NAVY, True)
    txt(slide, body, x + 0.22, y + 0.68, w - 0.42, h - 0.84, body_size, INK)


def note(slide, text, y=6.64, color=MUTED):
    txt(slide, text, 0.63, y, 12.0, 0.31, 9.5, color)


# 1 — executive story
s = base("01 / Bài toán", "VBot: trợ lý AI văn phòng, chi phí tháng thấp nhất", "Một cloud chính • 2.000 DAU • 5 cấp tác vụ • chất lượng và kiểm soát trước tối ưu chi phí")
pill(s, "BẢN NHÁP — CHƯA KHÓA GIÁ", 0.62, 1.91, 2.74, ORANGE)
card(s, 0.62, 2.42, 3.86, 2.86, "Vấn đề", "Tri thức phân tán giữa Google Workspace, Microsoft 365, Notion, Jira, Confluence; agent cần làm việc mà không vượt quyền.")
card(s, 4.74, 2.42, 3.86, 2.86, "Giải pháp", "AWS Singapore, truy cập theo danh tính; RAG lọc ACL, model routing, connector OAuth, sandbox cô lập và phê duyệt thao tác ghi.")
card(s, 8.86, 2.42, 3.86, 2.86, "Điều kiện thắng", "Giảm USD/tháng trong số phương án vượt H1–H8. Không dùng giá khuyến mãi; mỗi dòng giá có bằng chứng và khớp slide–XLSX.", ORANGE)
txt(s, "2.000", 0.75, 5.67, 2.4, 0.58, 30, TEAL, True)
txt(s, "người dùng/ngày", 0.75, 6.18, 2.7, 0.25, 12, MUTED)
txt(s, "19", 4.88, 5.67, 1.5, 0.58, 30, TEAL, True)
txt(s, "component bắt buộc", 4.88, 6.18, 2.6, 0.25, 12, MUTED)
txt(s, "99,5%", 9.01, 5.67, 2.7, 0.58, 30, TEAL, True)
txt(s, "mục tiêu giờ hành chính", 9.01, 6.18, 3.0, 0.25, 12, MUTED)
note(s, "Nguồn: vbot-capstone.html. Đây là thiết kế kiến trúc/báo giá, không phải hệ thống đã triển khai.")

# 2 — sizing
s = base("02 / Sizing", "Từ 2.000 DAU tới tải giờ cao điểm", "10 request/người/ngày và peak ×4 là giả định chung để lập dự toán, chưa được đo")
for x, value, label in [
    (0.64, "440.000", "request/tháng"), (3.72, "1.055,56M", "input token/tháng"),
    (6.80, "301,84M", "output token/tháng"), (9.88, "2,22", "request/giây peak")]:
    box(s, x, 1.95, 2.79, 1.23)
    txt(s, value, x + 0.16, 2.15, 2.5, 0.49, 23, TEAL, True)
    txt(s, label, x + 0.16, 2.72, 2.5, 0.23, 11, MUTED)
txt(s, "CƠ CẤU NĂM CẤP TÁC VỤ", 0.68, 3.57, 5.8, 0.31, 13, NAVY, True)
tiers = [("01  Email / hỏi đáp", 50), ("02  Tài liệu / slide", 25), ("03  Dữ liệu / code", 12), ("04  Agent dự án", 8), ("05  Thiết kế", 5)]
for i, (label, pct) in enumerate(tiers):
    y = 4.04 + i * 0.51
    txt(s, label, 0.71, y, 3.02, 0.28, 12, INK, True)
    box(s, 3.83, y + 0.03, 5.48, 0.18, LINE, LINE, False)
    box(s, 3.83, y + 0.03, 5.48 * pct / 50, 0.18, TEAL, TEAL, False)
    txt(s, f"{pct}%", 9.44, y - 0.01, 0.58, 0.30, 12, TEAL, True)
box(s, 10.34, 3.54, 2.31, 2.92, PALE, PALE)
txt(s, "HỆ QUẢ", 10.56, 3.78, 1.9, 0.3, 13, ORANGE, True)
txt(s, "59.840 sandbox run/tháng\n18,1 run đồng thời peak\n67 phiên chat peak\n220 giờ làm việc/tháng", 10.55, 4.24, 1.87, 1.92, 12, INK)
note(s, "Công thức: 2.000 × 22 × 10 = 440.000; peak = 440.000 / (220 × 3.600) × 4. Nguồn: dung/ASSUMPTIONS.md.")

# 3 — architecture, all 19 codes
s = base("03 / Kiến trúc", "Một kiến trúc AWS Singapore — đủ 19/19 component", "Prod hai AZ, staging tách biệt; Bedrock Nova text định tuyến trong APAC")
groups = [
    ("EDGE & QUYỀN", "authn   gateway   network   load_balancer   authz", 0.68, 1.94, 12.0, 0.75),
    ("ĐIỀU PHỐI", "compute   governance   cache_queue", 0.68, 2.91, 12.0, 0.75),
    ("MODEL & TÁC VỤ", "llm   embedding*   sandbox   connectors   secrets_kms", 0.68, 3.88, 12.0, 0.75),
    ("DỮ LIỆU & VẬN HÀNH", "database   vector_db   object_storage   backup_dr   monitoring   cicd   ops_labor", 0.68, 4.85, 12.0, 0.89),
]
for i, (label, codes, x, y, w, h) in enumerate(groups):
    box(s, x, y, w, h, WHITE)
    box(s, x, y, 1.77, h, TEAL if i < 3 else NAVY, TEAL if i < 3 else NAVY, False)
    txt(s, label, x + 0.11, y + 0.19, 1.52, 0.42, 10.5, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    txt(s, codes, x + 2.00, y + 0.13, w - 2.20, h - 0.20, 15 if i < 3 else 12.5, INK, True, valign=MSO_ANCHOR.MIDDLE)
for y in [2.73, 3.70, 4.67]:
    txt(s, "↓", 6.36, y, 0.5, 0.21, 17, ORANGE, True, PP_ALIGN.CENTER)
note(s, "* embedding là component tùy chọn ngoài 19 mã bắt buộc. API Gateway → VPC link → ALB → orchestrator; Bảo còn chốt SKU/HA.", 6.32)

# 4 — request sequence
s = base("04 / Luồng request", "Mỗi request đi qua quyền, policy và đường trả lời có kiểm soát", "Ví dụ: đọc tài liệu nội bộ rồi đề xuất cập nhật Jira")
steps = [
    ("01", "SSO + quota", "Xác thực phiên, tenant/user, hạn mức."),
    ("02", "RAG theo ACL", "Lọc vector; kiểm lại quyền ở nguồn."),
    ("03", "DLP + routing", "Che PII; chọn Micro/Pro sau eval."),
    ("04", "Tool hoặc connector", "Sandbox không token; ghi Jira cần duyệt."),
    ("05", "Audit + trả lời", "Kiểm output, ghi trace, stream kết quả."),
]
for i, (num, title, body) in enumerate(steps):
    x = 0.65 + i * 2.51
    box(s, x, 2.17, 2.22, 3.12, WHITE)
    pill(s, num, x + 0.16, 2.37, 0.53, ORANGE if i == 3 else TEAL)
    txt(s, title, x + 0.15, 2.90, 1.94, 0.56, 17, NAVY, True)
    txt(s, body, x + 0.15, 3.73, 1.94, 1.00, 13, INK)
    if i < 4:
        txt(s, "→", x + 2.21, 3.35, 0.31, 0.4, 20, ORANGE, True)
box(s, 0.68, 5.70, 12.0, 0.64, CYAN, CYAN)
txt(s, "Nguyên tắc: model chỉ đề xuất. Policy và người dùng quyết định hành động ghi; payload đổi thì xin duyệt lại.", 0.91, 5.86, 11.5, 0.35, 14, TEAL, True)
note(s, "Nguồn: dung/ARCHITECTURE.md và bao/SECURITY_DESIGN.md.")

# 5 — model route
s = base("05 / Model serving", "Route theo độ khó, không hạ chất lượng để giảm giá", "Nova Micro / Nova Pro qua Bedrock APAC; cấp 5 mặc định tạo HTML/SVG")
rows = [
    ("Cấp 1", "Email, hỏi đáp", "90% Micro / 10% Pro", "220.000"),
    ("Cấp 2", "Tài liệu, slide", "50% Micro / 50% Pro", "110.000"),
    ("Cấp 3", "Bảng tính, code", "Pro + sandbox", "52.800"),
    ("Cấp 4", "Agent nhiều bước", "Pro + queue + approval", "35.200"),
    ("Cấp 5", "Mockup, sơ đồ, minh họa", "Micro/Pro → HTML/SVG", "22.000"),
]
heads = [("CẤP", 0.87, 1.5), ("CÔNG VIỆC", 2.6, 3.6), ("ROUTING ĐỀ XUẤT", 6.63, 3.55), ("REQUEST/THÁNG", 10.44, 2.1)]
box(s, 0.65, 1.97, 12.02, 0.47, NAVY, NAVY)
for h, x, w in heads:
    txt(s, h, x, 2.08, w, 0.2, 10.5, WHITE, True)
for i, row in enumerate(rows):
    y = 2.46 + i * 0.67
    box(s, 0.65, y, 12.02, 0.63, WHITE if i % 2 == 0 else CYAN, WHITE if i % 2 == 0 else CYAN, False)
    for j, (value, x, w) in enumerate(zip(row, [0.87, 2.6, 6.63, 10.44], [1.5, 3.6, 3.55, 2.1])):
        txt(s, value, x, y + 0.17, w, 0.25, 13, TEAL if j == 0 else INK, j == 0 or j == 3)
box(s, 0.68, 6.10, 12.0, 0.5, PALE, PALE)
txt(s, "Ảnh bitmap ngoài APAC KHÔNG nằm trong baseline; chỉ bật sau duyệt residency và cộng riêng 4.400 ảnh / $176/tháng.", 0.85, 6.23, 11.7, 0.25, 12, ORANGE, True)
note(s, "Nguồn: kiet/KIET_MODEL_SERVING.md. Tỉ lệ route là giả định, chưa qua eval.")

# 6 — economics and quality
s = base("06 / Lựa chọn model", "API managed là baseline; tiết kiệm $774,95/tháng có điều kiện", "Chất lượng ≥95% so với Pro theo từng cấp; p95 TTFT ≤3 giây là mục tiêu chưa đo")
card(s, 0.68, 1.95, 3.71, 2.55, "Micro + Pro", "$1.669,00/tháng text\nChọn tạm; chỉ áp dụng tỉ lệ route nếu Micro vượt eval 100 case/cấp.", TEAL, 15)
card(s, 4.81, 1.95, 3.71, 2.55, "Toàn bộ Pro", "$2.443,95/tháng text\nDễ vận hành hơn, nhưng thêm $774,95/tháng theo workload giả định.", ORANGE, 15)
card(s, 8.94, 1.95, 3.71, 2.55, "Tự host / hybrid", "Chưa định giá GPU Singapore. Ngưỡng lý thuyết ≤$1,143/node-giờ cho 2 node; phải benchmark.", NAVY, 15)
box(s, 0.68, 4.85, 12.0, 1.29, CYAN, CYAN)
txt(s, "Cổng nghiệm thu", 0.91, 5.02, 2.2, 0.33, 16, TEAL, True)
txt(s, "≥95% điểm Pro/cấp • 0 lỗi ACL/PII • load test 2,22 RPS / 30 phút • agent ≤10 phút gồm queue/tool/sandbox", 3.04, 5.03, 9.3, 0.66, 13, INK)
note(s, "Giá phần text theo AWS Bedrock Price List Singapore; cache hit giả định 0%, không trừ tiền khi chưa đo. Nguồn: kiet/KIET_MODEL_SERVING.md.")

# 7 — data/Phong
s = base("07 / Data & workspace", "RAG chỉ trả tài liệu user được phép xem", "Phần này tổng hợp từ kế hoạch và giả định Dũng; chưa có tài liệu bàn giao riêng của Phong")
card(s, 0.68, 1.95, 3.72, 3.45, "Connectors + OAuth", "Google Workspace, Microsoft 365, Outlook, Notion, Jira, Confluence. Đọc theo scope user; ghi chỉ sau approval. Ước lượng 564.960 call/tháng.", TEAL, 14)
card(s, 4.82, 1.95, 3.72, 3.45, "RAG + ACL", "Index gắn nguồn, tenant, ACL version. Lọc trước truy hồi và kiểm lại tại nguồn; cache key mang quyền. Cohere Embed Multilingual v3 cho query tiếng Việt/Anh.", TEAL, 14)
card(s, 8.96, 1.95, 3.72, 3.45, "Workspace riêng", "Quota 10 GB/user = trần logic 20 TB; giả định dùng thực 1 GB/user ≈2 TB. Database, vector DB, object storage và queue tách ranh giới giá.", ORANGE, 14)
box(s, 0.68, 5.70, 12.0, 0.64, PALE, PALE)
txt(s, "Cần Phong xác nhận: corpus/GB thực, MAU/license, ACL gốc, rate limit connector, replica/index và giá 5 component.", 0.88, 5.86, 11.55, 0.35, 13, ORANGE, True)
note(s, "Nguồn: dung/ASSUMPTIONS.md, dung/ARCHITECTURE.md, TEAMMATE_PLAN.md. Số liệu data là giả định quy hoạch.")

# 8 — sandbox
s = base("08 / Sandbox", "Code chạy trong môi trường tạm — không mang token, không có egress", "59.840 run/tháng; peak ≈18,1 run đồng thời nếu 60 giây/run")
for i, (n, label, body) in enumerate([
    ("01", "Stage đúng quyền", "Runner nhận input đã lọc ACL; chỉ mount dữ liệu đúng user."),
    ("02", "Cô lập runtime", "Ứng viên gVisor/microVM; non-root, không đặc quyền, không cloud metadata."),
    ("03", "Giới hạn & hủy", "Cấm mạng mặc định; quota CPU/RAM/PID/disk/thời gian; dọn sau chạy."),
    ("04", "Kiểm artifact", "Scan output, DLP và publish qua runner vào workspace đúng chủ.")]):
    x = 0.68 + i * 3.08
    box(s, x, 2.20, 2.76, 3.36, WHITE)
    pill(s, n, x + 0.16, 2.42, 0.52)
    txt(s, label, x + 0.16, 3.01, 2.40, 0.55, 16, NAVY, True)
    txt(s, body, x + 0.16, 3.79, 2.4, 1.16, 13, INK)
box(s, 0.68, 5.81, 12.0, 0.52, PALE, PALE)
txt(s, "Chưa chốt runtime/SKU, vCPU/RAM và billing minimum; Bảo phải chứng minh cách chịu peak khi mất một AZ.", 0.89, 5.94, 11.54, 0.28, 12.5, ORANGE, True)
note(s, "Nguồn: bao/SLIDES_8_10.md và bao/SIZING_AND_COST.md.")

# 9 — security
s = base("09 / Governance", "Model đề xuất; policy và người dùng quyết định", "SSO/MFA, RBAC/ABAC, ACL gốc, token vault, DLP và audit 12 tháng")
rows = [
    ("DỮ LIỆU VÀO", "Connector/RAG là nguồn không tin cậy → phân loại, kiểm ACL, che PII/secret, chặn prompt injection."),
    ("HÀNH ĐỘNG GHI", "Preview đích + payload → user xác nhận → kiểm lại scope/quyền → connector thực thi idempotent."),
    ("TOKEN & SANDBOX", "OAuth ủy nhiệm nằm trong vault KMS; model/sandbox không nhận token; mặc định deny egress."),
    ("DỮ LIỆU RA", "Guardrail/DLP kiểm kết quả và artifact; audit 12 tháng; đổi model/prompt qua eval gate."),
]
for i, (label, body) in enumerate(rows):
    y = 2.0 + i * 1.07
    box(s, 0.68, y, 12.0, 0.87, WHITE)
    box(s, 0.68, y, 2.25, 0.87, TEAL if i != 1 else ORANGE, TEAL if i != 1 else ORANGE, False)
    txt(s, label, 0.88, y + 0.29, 1.89, 0.27, 11.5, WHITE, True)
    txt(s, body, 3.17, y + 0.19, 9.17, 0.51, 13.5, INK)
note(s, "Nguồn: bao/SECURITY_DESIGN.md. Cần kiểm điều khoản xử lý dữ liệu của Bedrock APAC trước khi tuyên bố đáp ứng H1/H5.")

# 10 — HA/ops
s = base("10 / HA & vận hành", "Thiết kế cho giờ hành chính; chỉ công bố đạt sau diễn tập", "99,5% availability • TTFT p95 ≤3s • agent ≤10 phút • RTO 4h / RPO 1h")
card(s, 0.68, 2.0, 3.72, 3.2, "Hai AZ + staging", "API Gateway/ALB managed; orchestrator hai replica; DB failover; staging tách quyền/dữ liệu. AZ còn lại cần đủ tải hoặc policy giảm tải đã duyệt.", TEAL, 14)
card(s, 4.82, 2.0, 3.72, 3.2, "Backup + restore", "PITR/backup object, kiểm bản sao và khóa; restore drill đo RTO 4 giờ, RPO 1 giờ. Hai AZ không đồng nghĩa DR khi mất cả region.", ORANGE, 14)
card(s, 8.96, 2.0, 3.72, 3.2, "Quan sát + FinOps", "Trace TTFT/queue/tool, token/user, sandbox, backup age. Cảnh báo 70/90/100% ngân sách; CI/CD staging → eval → canary → rollback.", NAVY, 14)
box(s, 0.68, 5.62, 12.0, 0.73, CYAN, CYAN)
txt(s, "Trạng thái: chưa có kết quả load test, failover test hay restore drill; các SLO trên là mục tiêu nghiệm thu.", 0.89, 5.85, 11.46, 0.32, 13, TEAL, True)
note(s, "Nguồn: bao/OPERATIONS_AND_ACCEPTANCE.md và dung/ARCHITECTURE.md.")

# 11 — cost incomplete
s = base("11 / Chi phí", "Chưa thể khóa TOTAL: hiện chỉ có phần Kiệt và nhân sự Dũng", "Không cộng giả các component chưa có SKU, đơn giá và bằng chứng")
card(s, 0.68, 2.02, 3.72, 2.38, "Model + embedding + gateway", "$1.679,86 / tháng\nBaseline không bitmap; giá gateway vẫn cần đối chiếu SKU Singapore.", TEAL, 15)
card(s, 4.82, 2.02, 3.72, 2.38, "Nhân sự vận hành", "$800,00 / tháng\nGiả định 40 giờ × $20; Dũng cần xác nhận ca trực và đơn giá.", ORANGE, 15)
card(s, 8.96, 2.02, 3.72, 2.38, "Đã có số sơ bộ", "$2.479,86 / tháng\nKHÔNG PHẢI TOTAL: thiếu data, sandbox, HA, network, backup và bảo mật.", RED, 15)
box(s, 0.68, 4.72, 12.0, 1.3, PALE, PALE)
txt(s, "Muốn khóa slide này", 0.93, 4.96, 2.52, 0.39, 16, ORANGE, True)
txt(s, "Nhận 5 component của Phong + 11 component của Bảo → ghép BaoGia → xác minh bằng chứng → lấy TOTAL, USD/DAU và top 5 trực tiếp từ XLSX.", 3.42, 4.94, 8.82, 0.70, 13, INK)
note(s, "Đề yêu cầu TOTAL và top 5; slide này cố ý đánh dấu chưa khóa, không dùng làm bản nộp cuối. Nguồn: kiet/KIET_BAOGIA.tsv, dung/DUNG_BAOGIA.tsv.", 6.44, RED)

# 12 — growth/gates
s = base("12 / Mở rộng & quyết định", "4.000 DAU: tải gần gấp đôi, chi phí không tự động gấp đôi", "Chốt thiết kế sau khi đủ giá và vượt cổng chất lượng H1–H8")
for i, (value, label) in enumerate([
    ("880.000", "request/tháng"), ("4,44", "request/giây peak"),
    ("119.680", "sandbox run/tháng"), ("$3.359,20", "phần Kiệt/tháng*")]):
    x = 0.68 + i * 3.07
    box(s, x, 2.0, 2.76, 1.33)
    txt(s, value, x + 0.13, 2.22, 2.50, 0.47, 23, TEAL, True)
    txt(s, label, x + 0.14, 2.82, 2.46, 0.25, 11, MUTED)
card(s, 0.68, 3.71, 5.79, 2.0, "Rủi ro cần xử lý", "Geo inference/residency; ACL stale; sandbox peak khi mất AZ; thiếu bằng chứng giá; TTFT và eval chưa đo.", ORANGE, 14)
card(s, 6.90, 3.71, 5.79, 2.0, "Cổng hoàn thành", "Chốt SKU và TOTAL; eval ≥95% từng cấp; load/HA/restore test; đối chiếu PDF với XLSX và chạy validator ĐẠT.", TEAL, 14)
note(s, "* Giữ mix và đơn giá; corpus index không nhân đôi. Tổng hệ thống phải tính lại từng meter, replica, staging và license.")

prs.save(PPTX)
print(PPTX)
